from __future__ import annotations

import argparse
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from src.pipeline import ResumeScreeningPipeline


logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s"
)


def get_decision(score: float) -> tuple[str, str]:

    if score >= 0.70:
        return "SELECTED", "Strong overall match with the job description."

    elif score >= 0.60:
        return "REVIEW", "Moderate match. Recruiter should review the candidate."

    else:
        return "REJECTED", "Low overall match with the job description."


def create_excel_report(ranked, output_file):

    workbook = Workbook()
    sheet = workbook.active

    sheet.title = "Screening Results"

    headers = [
        "Rank",
        "Resume File",
        "Overall Match Score",
        "Semantic Match",
        "Skill Match",
        "Decision",
        "Reason",
        "Matched Skills",
        "Missing Skills",
        "Experience",
        "Education"
    ]

    sheet.append(headers)

    # Header formatting
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Add candidate data
    for rank, candidate in enumerate(ranked, start=1):

        decision, reason = get_decision(
            candidate.final_score
        )

        row = [
            rank,
            candidate.resume_id,
            round(candidate.final_score * 100, 2),
            round(candidate.semantic_score * 100, 2),
            round(candidate.skill_overlap_score * 100, 2),
            decision,
            reason,
            ", ".join(candidate.matched_skills)
            if candidate.matched_skills
            else "None",
            ", ".join(candidate.missing_skills)
            if candidate.missing_skills
            else "None",
            candidate.years_experience,
            ", ".join(candidate.education)
            if candidate.education
            else "None"
        ]

        sheet.append(row)

    # Format all cells
    for row in sheet.iter_rows():

        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    # Set column widths
    column_widths = {
        1: 8,
        2: 25,
        3: 20,
        4: 18,
        5: 15,
        6: 15,
        7: 45,
        8: 40,
        9: 40,
        10: 15,
        11: 20
    }

    for column, width in column_widths.items():
        sheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    # Freeze header row
    sheet.freeze_panes = "A2"

    # Enable filtering
    sheet.auto_filter.ref = sheet.dimensions

    workbook.save(output_file)

    print(
        f"\nExcel report created successfully: {output_file}"
    )


def main():

    parser = argparse.ArgumentParser(
        description="Resume Matching and Candidate Selection System"
    )

    parser.add_argument(
        "--resumes_dir",
        required=True,
        help="Directory containing resume files (pdf/docx/txt)"
    )

    parser.add_argument(
        "--job_description",
        required=True,
        help="Path to a .txt file containing the job description"
    )

    parser.add_argument(
        "--top_k",
        type=int,
        default=10,
        help="Number of candidates to display"
    )

    parser.add_argument(
        "--min_similarity",
        type=float,
        default=0.0
    )

    parser.add_argument(
        "--min_years_experience",
        type=float,
        default=None
    )

    parser.add_argument(
        "--min_education",
        type=str,
        default=None,
        choices=[
            None,
            "High School",
            "Associate",
            "Bachelors",
            "Masters",
            "PhD"
        ]
    )

    parser.add_argument(
        "--save_index_to",
        type=str,
        default=None,
        help="Optional directory to save the FAISS index"
    )

    args = parser.parse_args()

    jd_text = Path(
        args.job_description
    ).read_text(
        encoding="utf-8",
        errors="ignore"
    )

    pipeline = ResumeScreeningPipeline(
        batch_size=128
    )

    print("\n" + "=" * 70)
    print("          RESUME MATCHING & SELECTION SYSTEM")
    print("=" * 70)

    print(
        f"\nIndexing resumes from: {args.resumes_dir}"
    )

    index_report = pipeline.index_directory(
        args.resumes_dir,
        batch_size=150
    )

    print(
        f"Indexed: {index_report['indexed']}  "
        f"Failed: {len(index_report['failed'])}"
    )

    print(
        f"Performance: {index_report['performance']}"
    )

    if index_report["failed"]:

        print("\nFailed resumes:")

        for resume_id in index_report["failed"]:
            print(f"  - {resume_id}")

    if args.save_index_to:

        pipeline.save_index(
            args.save_index_to
        )

        print(
            f"\nIndex saved to: {args.save_index_to}"
        )

    ranked = pipeline.match(
        job_description_text=jd_text,
        top_k=args.top_k,
        min_similarity=args.min_similarity,
        min_years_experience=args.min_years_experience,
        min_education=args.min_education,
    )

    # ---------------- TERMINAL RESULTS ----------------

    print("\n" + "=" * 70)
    print("                    SCREENING RESULTS")
    print("=" * 70)

    print(
        f"\nCandidates analyzed: {index_report['indexed']}"
    )

    print(
        f"Candidates displayed: {len(ranked)}"
    )

    print("\n" + "-" * 70)

    print(
        f"{'Rank':<6}"
        f"{'Candidate':<20}"
        f"{'Score':<12}"
        f"{'Decision'}"
    )

    print("-" * 70)

    for i, candidate in enumerate(
        ranked,
        start=1
    ):

        decision, reason = get_decision(
            candidate.final_score
        )

        score_percentage = (
            candidate.final_score * 100
        )

        print(
            f"{i:<6}"
            f"{candidate.resume_id:<20}"
            f"{score_percentage:>6.1f}%     "
            f"{decision}"
        )

    print("-" * 70)

    # ---------------- EXCEL REPORT ----------------

    output_file = Path(
        "resume_screening_results.xlsx"
    )

    create_excel_report(
        ranked,
        output_file
    )

    print("\n" + "=" * 70)
    print("                  SCREENING COMPLETE")
    print("=" * 70)


if __name__ == "__main__":
    main()